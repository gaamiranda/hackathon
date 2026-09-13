"""Local, deterministic document → plain text. Nothing binary ever reaches the LLM (PLAN.md risk 2)."""

import io
from pathlib import PurePosixPath

from openpyxl import load_workbook
from pypdf import PdfReader

from procureai.domain.models import QuoteSource

EXTENSIONS = {
    ".pdf": QuoteSource.PDF,
    ".xlsx": QuoteSource.XLSX,
    ".xlsm": QuoteSource.XLSX,
    ".txt": QuoteSource.EMAIL,
    ".eml": QuoteSource.EMAIL,
    ".md": QuoteSource.EMAIL,
}


class UnsupportedDocument(ValueError):
    pass


def detect_source(filename: str) -> QuoteSource:
    ext = PurePosixPath(filename).suffix.lower()
    try:
        return EXTENSIONS[ext]
    except KeyError:
        raise UnsupportedDocument(f"unsupported file type {ext or '(none)'}; use {sorted(EXTENSIONS)}") from None


def pdf_to_text(data: bytes) -> str:
    return "\n".join((page.extract_text() or "").strip() for page in PdfReader(io.BytesIO(data)).pages).strip()


def xlsx_to_text(data: bytes) -> str:
    """Cached values only (data_only=True). Two-cell rows become "label: value";
    wider rows are joined with " | ". Sheets are separated by "## <name>" headers."""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        lines: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(v) for v in row if v is not None and str(v).strip() != ""]
            if not cells:
                continue
            lines.append(f"{cells[0]}: {cells[1]}" if len(cells) == 2 else " | ".join(cells))
        if lines:
            blocks.append((f"## {ws.title}\n" if len(wb.worksheets) > 1 else "") + "\n".join(lines))
    return "\n\n".join(blocks).strip()


def _cell(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def to_text(filename: str, data: bytes) -> tuple[QuoteSource, str]:
    source = detect_source(filename)
    if source is QuoteSource.PDF:
        return source, pdf_to_text(data)
    if source is QuoteSource.XLSX:
        return source, xlsx_to_text(data)
    return source, data.decode("utf-8", errors="replace").strip()
